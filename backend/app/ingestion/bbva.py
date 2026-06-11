"""BBVA Spain `.xlsx` export adapter.

Layout (sheet "Informe BBVA", but the sheet name is not assumed): a few
preamble rows, then a header row
    F.Valor | Fecha | Concepto | Movimiento | Importe | Divisa | Disponible | Divisa | Observaciones
possibly preceded by an empty leading column. Columns are mapped by header
name, never by position.
"""

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from .base import (
    NormalizedTransaction,
    StatementFormatError,
    cell_text,
    parse_amount_cents,
    parse_date,
)

# Headers that must all appear on one row for it to be the header row.
_REQUIRED = ("fecha", "concepto", "importe")


class BBVAXlsxAdapter:
    source = "import_bbva"

    def parse(self, content: bytes) -> list[NormalizedTransaction]:
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise StatementFormatError() from exc
        try:
            for sheet in workbook.worksheets:
                rows = [list(r) for r in sheet.iter_rows(values_only=True)]
                parsed = self._parse_rows(rows)
                if parsed is not None:
                    return parsed
        finally:
            workbook.close()
        raise StatementFormatError()

    def _parse_rows(self, rows: list[list[Any]]) -> list[NormalizedTransaction] | None:
        header_index, columns = _find_header(rows)
        if columns is None:
            return None
        transactions: list[NormalizedTransaction] = []
        for row in rows[header_index + 1 :]:
            cells = {name: row[i] if i < len(row) else None for name, i in columns.items()}
            if not cell_text(cells.get("fecha")) and not cell_text(cells.get("importe")):
                continue  # trailing blank/footer rows
            concept = cell_text(cells.get("concepto"))
            movement = cell_text(cells.get("movimiento"))
            description = f"{concept} — {movement}" if movement and movement != concept else concept
            balance_cents = _optional_cents(cells.get("disponible"))
            transactions.append(
                NormalizedTransaction(
                    date=parse_date(cells.get("fecha")),
                    amount_cents=parse_amount_cents(cells.get("importe")),
                    currency=cell_text(cells.get("divisa")) or "EUR",
                    description=description,
                    # No per-row id in the export; the running balance keeps
                    # legitimate same-day identical rows distinguishable (D1/D3).
                    external_ref=f"{description}|bal:{balance_cents}"
                    if balance_cents is not None
                    else None,
                    raw_payload={k: cell_text(v) for k, v in cells.items() if v is not None},
                )
            )
        return transactions


def _find_header(rows: list[list[Any]]) -> tuple[int, dict[str, int] | None]:
    for index, row in enumerate(rows):
        names = [cell_text(c).lower() for c in row]
        if all(any(n.startswith(req) for n in names) for req in _REQUIRED):
            columns: dict[str, int] = {}
            for position, name in enumerate(names):
                if name.startswith("f.valor"):
                    columns["f_valor"] = position
                elif name == "fecha":
                    columns["fecha"] = position
                elif name.startswith("concepto"):
                    columns["concepto"] = position
                elif name.startswith("movimiento"):
                    columns["movimiento"] = position
                elif name.startswith("importe"):
                    columns["importe"] = position
                elif name.startswith("divisa") and "divisa" not in columns:
                    columns["divisa"] = position
                elif name.startswith("disponible"):
                    columns["disponible"] = position
                elif name.startswith("observaciones"):
                    columns["observaciones"] = position
            if {"fecha", "concepto", "importe"} <= columns.keys():
                return index, columns
    return -1, None


def _optional_cents(value: Any) -> int | None:
    if value is None or cell_text(value) == "":
        return None
    try:
        return parse_amount_cents(value)
    except StatementFormatError:
        return None
